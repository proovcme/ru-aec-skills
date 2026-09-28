#!/usr/bin/env python3
"""Build the public neutral RIM LSR workbook template."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EAF3F8"
GRAY = "E7E6E6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="7F8C8D")


def merge_value(ws, cell_range: str, value: str, *, bold: bool = False, fill: str | None = None) -> None:
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":", 1)[0]]
    cell.value = value
    cell.font = Font(name="Arial", size=10, bold=bold, color=WHITE if fill == BLUE else "000000")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ЛСР РИМ"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = None

    widths = [7, 22, 58, 14, 14, 12, 16, 18, 11, 18, 12, 19]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    merge_value(ws, "A1:L1", "Форма локального сметного расчёта (сметы) для ресурсно-индексного метода", bold=True, fill=BLUE)
    ws.row_dimensions[1].height = 28
    requisites = [
        (3, "Наименование программного продукта", "<программный продукт и версия>"),
        (4, "Наименование редакции сметных нормативов", "<редакция ФСНБ>"),
        (5, "Реквизиты приказа Минстроя России об утверждении дополнений и изменений к сметным нормативам", "<дата и номер приказа>"),
        (6, "Реквизиты письма об индексах", "<дата, номер, приложение>"),
        (7, "Реквизиты акта об оплате труда", "<акт субъекта РФ>"),
        (8, "Обоснование текущих цен", "<ФГИС ЦС / КА / иной источник>"),
        (9, "Субъект Российской Федерации", "<субъект>"),
        (10, "Ценовая зона", "<зона>"),
        (11, "Наименование стройки", "<стройка>"),
        (12, "Объект капитального строительства", "<объект>"),
    ]
    for row, label, value in requisites:
        merge_value(ws, f"A{row}:D{row}", label, bold=True, fill=PALE_BLUE)
        merge_value(ws, f"E{row}:L{row}", value)
        ws.row_dimensions[row].height = 30 if row == 5 else 22

    merge_value(ws, "A14:L14", "ЛОКАЛЬНЫЙ СМЕТНЫЙ РАСЧЁТ (СМЕТА) № <номер>", bold=True, fill=BLUE)
    merge_value(ws, "A15:L15", "<наименование работ и затрат>", bold=True)
    merge_value(ws, "A16:L16", "Составлен ресурсно-индексным методом")
    merge_value(ws, "A18:D18", "Основание", bold=True, fill=PALE_BLUE)
    merge_value(ws, "E18:L18", "<проектная документация; ВОР>")
    merge_value(ws, "A19:D19", "Текущий уровень цен", bold=True, fill=PALE_BLUE)
    merge_value(ws, "E19:L19", "<квартал и год>")
    merge_value(ws, "A20:D20", "Статус расчёта", bold=True, fill=PALE_BLUE)
    merge_value(ws, "E20:L20", "<подтверждено / предварительно / требует уточнения>")

    summary = [
        (21, "Сметная стоимость, тыс. руб.", "Средства на оплату труда рабочих, тыс. руб."),
        (22, "в том числе:", "Средства на оплату труда машинистов, тыс. руб."),
        (23, "строительных работ, тыс. руб.", "Нормативные затраты труда рабочих, чел.-ч"),
        (24, "монтажных работ, тыс. руб.", "Нормативные затраты труда машинистов, чел.-ч"),
        (25, "оборудования, тыс. руб.", ""),
        (26, "прочих затрат, тыс. руб.", ""),
    ]
    for row, left, right in summary:
        merge_value(ws, f"A{row}:D{row}", left, bold=row == 21, fill=PALE_BLUE if row == 21 else None)
        ws[f"E{row}"] = "<значение>"
        merge_value(ws, f"G{row}:J{row}", right, bold=row == 21, fill=PALE_BLUE if row == 21 else None)
        if right:
            ws[f"K{row}"] = "<значение>"

    headers1 = ["№ п/п", "Обоснование", "Наименование работ и затрат", "Единица измерения", "Количество", "", "", "Сметная стоимость, руб.", "", "", "", ""]
    headers2 = ["", "", "", "", "на единицу измерения", "коэффициенты", "всего с учётом коэффициентов", "на единицу в базисном уровне цен", "индекс", "на единицу в текущем уровне цен", "коэффициенты", "всего в текущем уровне цен"]
    for col, value in enumerate(headers1, start=1):
        ws.cell(28, col, value)
    for col, value in enumerate(headers2, start=1):
        ws.cell(29, col, value)
    for cell_range in ("A28:A29", "B28:B29", "C28:C29", "D28:D29", "E28:G28", "H28:L28"):
        ws.merge_cells(cell_range)
    for row in (28, 29):
        for col in range(1, 13):
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=BLUE)
            cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[28].height = 28
    ws.row_dimensions[29].height = 48

    merge_value(ws, "A31:L31", "Раздел 1. <наименование раздела>", bold=True, fill=LIGHT_BLUE)
    position_rows = {
        32: ["<№>", "<код нормы>", "<наименование нормы; ВОР-ID>", "<ед.>", "<объём>", 1, "<объём>", "", "", "", 1, "<итог>"],
        33: ["", "1", "ОТ (ЗТ)", "чел.-ч", "", "", "<ЗТ>", "", "", "", "", "<ОТ>"],
        34: ["", "<код труда>", "<средний разряд работы>", "чел.-ч", "<норма>", 1, "<ЗТ>", "", "", "<ставка>", 1, "<ОТ>"],
        35: ["", "2", "ЭМ", "", "", "", "", "", "", "", "", "<ЭМ>"],
        36: ["", "<код труда машиниста>", "в том числе ОТм (ЗТм)", "чел.-ч", "", "", "<ЗТм>", "", "", "<ставка>", 1, "<ОТм>"],
        37: ["", "<код машины>", "<машина или механизм>", "маш.-ч", "<норма>", 1, "<количество>", "<база>", "<индекс>", "<текущая>", 1, "<стоимость>"],
        38: ["", "4", "М", "", "", "", "", "", "", "", "", "<М>"],
        39: ["", "<код ресурса>", "<материальный ресурс>", "<ед.>", "<норма>", 1, "<количество>", "<база>", "<индекс>", "<текущая>", 1, "<стоимость>"],
        40: ["", "", "ЗТ", "чел.-ч", "", "", "<ЗТ>", "", "", "", "", ""],
        41: ["", "", "ЗТм", "чел.-ч", "", "", "<ЗТм>", "", "", "", "", ""],
        42: ["", "", "Итого прямые затраты", "", "", "", "", "", "", "", "", "<ПЗ>"],
        43: ["", "", "ФОТ", "", "", "", "", "", "", "", "", "<ОТ+ОТм>"],
        44: ["", "<пункт 812/пр>", "Накладные расходы", "%", "<норматив>", 1, "<итоговый %>", "", "", "", "", "<НР>"],
        45: ["", "<пункт 774/пр>", "Сметная прибыль", "%", "<норматив>", 1, "<итоговый %>", "", "", "", "", "<СП>"],
        46: ["", "", "Всего по позиции", "", "", "", "", "", "", "", "", "<итог>"],
    }
    for row, values in position_rows.items():
        for col, value in enumerate(values, start=1):
            ws.cell(row, col, value)

    totals = [
        (48, "Итого прямые затраты по разделу"), (49, "в том числе"), (50, "оплата труда (ОТ)"),
        (51, "эксплуатация машин и механизмов"), (52, "оплата труда машинистов (ОТм)"),
        (53, "материальные ресурсы"), (54, "перевозка"), (55, "Итого ФОТ"),
        (56, "Итого накладные расходы"), (57, "Итого сметная прибыль"),
        (58, "Итого оборудование"), (59, "Итого прочие затраты"), (60, "Итого по разделу"),
        (62, "Справочно"), (63, "материальные ресурсы, отсутствующие в ФРСН"),
        (64, "оборудование, отсутствующее в ФРСН"), (65, "затраты труда рабочих"),
        (66, "затраты труда машинистов"), (68, "ИТОГИ ПО СМЕТЕ"),
        (69, "ВСЕГО строительные / монтажные работы"), (70, "в том числе всего прямые затраты"),
        (71, "в том числе"), (72, "оплата труда (ОТ)"), (73, "эксплуатация машин и механизмов"),
        (74, "оплата труда машинистов (ОТм)"), (75, "материальные ресурсы"), (76, "перевозка"),
        (77, "Итого ФОТ"), (78, "Итого накладные расходы"), (79, "Итого сметная прибыль"),
        (80, "Итого оборудование"), (81, "Итого прочие затраты"), (82, "ВСЕГО ПО СМЕТЕ"),
    ]
    for row, label in totals:
        merge_value(ws, f"A{row}:K{row}", label, bold=row in {48, 60, 62, 68, 69, 82}, fill=LIGHT_BLUE if row in {48, 60, 62, 68, 69, 82} else None)
        ws[f"L{row}"] = "<значение>" if row not in {49, 62, 68, 71} else ""
    merge_value(ws, "A84:F84", "Составил ____________________")
    merge_value(ws, "G84:L84", "Проверил ____________________")

    for row in range(28, 83):
        for col in range(1, 13):
            cell = ws.cell(row, col)
            cell.font = Font(name="Arial", size=9, bold=cell.font.bold)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if col in range(5, 13):
                cell.number_format = '#,##0.00'

    audit = wb.create_sheet("Обоснования ЛСР")
    audit.sheet_view.showGridLines = False
    audit.freeze_panes = None
    headers = [
        "LSR-ID", "ВОР-ID", "Раздел / зона", "Код нормы / ресурса", "Редакция / GUID", "Ед. нормы",
        "Объём ВОР", "Объём нормы", "Коэффициент и пункт", "Источник цены", "Субъект / зона / период",
        "Строка ЛСР", "Доказательство применимости", "Статус", "Ответ / уточнение пользователя",
        "Решение после ответа", "price_kind", "price_base", "index", "Группа однородных ресурсов",
        "resolved_price", "zone_id", "period_id", "dataset / source row", "provenance",
    ]
    for col, value in enumerate(headers, start=1):
        cell = audit.cell(1, col, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        audit.column_dimensions[get_column_letter(col)].width = 18 if col not in {13, 24, 25} else 42
    audit.row_dimensions[1].height = 45

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.output)


if __name__ == "__main__":
    main()
