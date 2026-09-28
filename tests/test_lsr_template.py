import importlib.util
from pathlib import Path

from openpyxl import load_workbook


PATH = Path(__file__).parents[1] / "prepare-local-cost-estimate" / "scripts" / "build_lsr_template.py"
SPEC = importlib.util.spec_from_file_location("build_lsr_template", PATH)
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


def test_template_contains_full_rim_shell(tmp_path):
    output = tmp_path / "lsr-rim-template.xlsx"
    MODULE.build(output)
    workbook = load_workbook(output, data_only=False)
    sheet = workbook["ЛСР РИМ"]

    assert sheet.max_column == 12
    assert sheet.freeze_panes is None
    assert sheet["A5"].value.startswith("Реквизиты приказа Минстроя России")
    assert sheet["A68"].value == "ИТОГИ ПО СМЕТЕ"
    assert sheet["A82"].value == "ВСЕГО ПО СМЕТЕ"
    assert workbook["Обоснования ЛСР"].max_column == 25
